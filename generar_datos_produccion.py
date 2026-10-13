"""Genera datos SIMULADOS de producción de una hilandería (Abril-Septiembre 2026).
Todos los datos son inventados: no pertenecen a ninguna empresa real.
Salida: produccion_planta.xlsx  (hojas: Maquinas, Produccion, Paradas)"""
import random
from datetime import date, timedelta
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

random.seed(7)
np.random.seed(7)

MIN_TURNO = 480  # 8 horas

maquinas = pd.DataFrame([
    ("M01", "Carda 01", "Cardado", 180), ("M02", "Carda 02", "Cardado", 180),
    ("M03", "Manuar 01", "Estirado", 220), ("M04", "Manuar 02", "Estirado", 220),
    ("M05", "Open-End 01", "Hilado", 150), ("M06", "Open-End 02", "Hilado", 150),
    ("M07", "Continua 01", "Hilado", 130), ("M08", "Bobinadora 01", "Bobinado", 110),
], columns=["ID_Maquina", "Maquina", "Sector", "Meta_Kg_Turno"])

turnos = ["Mañana", "Tarde", "Noche"]
causas = ["Falla mecánica", "Falla eléctrica", "Cambio de producto",
          "Falta de materia prima", "Mantenimiento programado", "Limpieza", "Corte de energía"]
pesos_causas = [0.28, 0.12, 0.18, 0.12, 0.12, 0.12, 0.06]

inicio, fin = date(2026, 4, 1), date(2026, 9, 30)
dias = [inicio + timedelta(days=i) for i in range((fin - inicio).days + 1)]
dias = [d for d in dias if d.weekday() != 6]          # sin domingos

prod, paradas = [], []
for d in dias:
    mejora = 1 + 0.04 * (d - inicio).days / 183         # mejora gradual de eficiencia
    corte_energia = d in (date(2026, 7, 14), date(2026, 7, 15))
    for _, m in maquinas.iterrows():
        for t in turnos:
            # Cantidad de paradas del turno (Open-End 02 y turno Noche fallan más)
            lam = 1.4
            if m.Maquina == "Open-End 02": lam += 1.2
            if t == "Noche": lam += 0.5
            n_par = np.random.poisson(lam)
            total_par = 0
            for _ in range(n_par):
                causa = random.choices(causas, pesos_causas)[0]
                if m.Maquina == "Open-End 02" and random.random() < 0.4:
                    causa = "Falla mecánica"
                dur = int(np.clip(np.random.gamma(2.2, 14), 5, 150))
                total_par += dur
                paradas.append((d, t, m.ID_Maquina, causa, dur))
            if corte_energia:
                paradas.append((d, t, m.ID_Maquina, "Corte de energía", 180))
                total_par += 180
            total_par = min(total_par, MIN_TURNO - 60)
            efic = np.random.normal(0.93, 0.03) * mejora * (0.97 if t == "Noche" else 1.0)
            kg = m.Meta_Kg_Turno * (MIN_TURNO - total_par) / MIN_TURNO * min(efic, 1.02)
            desp_pct = np.clip(np.random.normal(0.035, 0.008)
                               + (0.015 if m.Maquina == "Open-End 02" else 0)
                               + (0.006 if t == "Noche" else 0), 0.01, 0.12)
            prod.append((d, t, m.ID_Maquina, MIN_TURNO, total_par,
                         round(kg, 1), m.Meta_Kg_Turno, round(kg * desp_pct, 1)))

df_prod = pd.DataFrame(prod, columns=["Fecha", "Turno", "ID_Maquina", "Minutos_Programados",
                                       "Minutos_Parada", "Kg_Producidos", "Meta_Kg", "Kg_Desperdicio"])
df_par = pd.DataFrame(paradas, columns=["Fecha", "Turno", "ID_Maquina", "Causa", "Minutos"])
# Los minutos de parada de Produccion se recalculan para que coincidan con la hoja Paradas
chk = df_par.groupby(["Fecha", "Turno", "ID_Maquina"])["Minutos"].sum().reset_index()
df_prod = df_prod.merge(chk, on=["Fecha", "Turno", "ID_Maquina"], how="left")
df_prod["Minutos_Parada"] = df_prod["Minutos"].fillna(0).clip(upper=MIN_TURNO - 60).astype(int)
df_prod = df_prod.drop(columns="Minutos")

with pd.ExcelWriter("produccion_planta.xlsx", engine="openpyxl", datetime_format="DD/MM/YYYY") as w:
    for nombre, datos in {"Maquinas": maquinas, "Produccion": df_prod, "Paradas": df_par}.items():
        datos.to_excel(w, sheet_name=nombre, index=False)
        ws = w.sheets[nombre]
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")
        for i, col in enumerate(datos.columns, 1):
            ancho = max(len(str(col)), *(len(str(v)) for v in datos[col].head(200)))
            ws.column_dimensions[get_column_letter(i)].width = min(ancho + 3, 30)
            if col == "Fecha":
                for c in ws[get_column_letter(i)][1:]:
                    c.number_format = "DD/MM/YYYY"
        ws.freeze_panes = "A2"

print(f"Produccion: {len(df_prod)} filas | Paradas: {len(df_par)} filas | Maquinas: {len(maquinas)}")
